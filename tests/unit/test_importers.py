"""In-memory integration fixtures for the supported resume upload formats."""

from io import BytesIO
from pathlib import Path
import sys
import unittest
from zipfile import ZIP_DEFLATED, ZipFile



from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.integrations.document_parsers.resume import (
    MAX_DOCX_UNCOMPRESSED_BYTES,
    MAX_TEXT_CHARS,
    MAX_UPLOAD_BYTES,
    extract_resume,
)


def pdf_bytes(text: str | None = None, *, pages: int = 1, password: str | None = None) -> bytes:
    writer = PdfWriter()
    for index in range(pages):
        page = writer.add_blank_page(width=612, height=792)
        if text is not None and index == 0:
            font = DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            })
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
            })
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(stream)
    if password is not None:
        writer.encrypt(password)
    result = BytesIO()
    writer.write(result)
    return result.getvalue()


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("김개발의 이력서")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "기술"
    table.cell(0, 1).text = "Python"
    table.cell(1, 0).text = "경력"
    table.cell(1, 1).text = "데이터 분석 3년"
    document.add_paragraph("프로젝트 성과: 처리 시간 30% 단축")
    result = BytesIO()
    document.save(result)
    return result.getvalue()


class TextImportTests(unittest.TestCase):
    def test_utf8_bom_and_cp949_korean(self):
        resume = "김개발\r\nPython 개발 경력 3년\r\n"
        for encoding in ("utf-8-sig", "cp949"):
            with self.subTest(encoding=encoding):
                self.assertEqual(extract_resume("이력서.TXT", resume.encode(encoding)), "김개발\nPython 개발 경력 3년")

    def test_binary_and_non_txt_contents_are_rejected(self):
        for data in (b"hello\x00world", b"a\x01\x02\x03\x04", b"%PDF-1.7\nPDF bytes", b"PK\x03\x04test"):
            with self.subTest(data=data):
                with self.assertRaises(ValueError):
                    extract_resume("resume.txt", data)

    def test_invalid_encoding_gives_actionable_error(self):
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            extract_resume("resume.txt", b"\xff\xff\xff")

    def test_empty_file_or_text(self):
        for data in (b"", b" \n\t "):
            with self.assertRaises(ValueError):
                extract_resume("resume.txt", data)

    def test_extension_and_upload_size_limit(self):
        with self.assertRaisesRegex(ValueError, "TXT, PDF, DOCX"):
            extract_resume("resume.exe", b"hello")
        with self.assertRaisesRegex(ValueError, "10MB"):
            extract_resume("resume.txt", b"a" * (MAX_UPLOAD_BYTES + 1))

    def test_text_limit_never_silently_truncates(self):
        self.assertEqual(len(extract_resume("resume.txt", b"a" * MAX_TEXT_CHARS)), MAX_TEXT_CHARS)
        with self.assertRaisesRegex(ValueError, "50,000"):
            extract_resume("resume.txt", b"a" * (MAX_TEXT_CHARS + 1))


class PdfImportTests(unittest.TestCase):
    def test_pdf_extracts_actual_text(self):
        self.assertIn("Python developer 3 years", extract_resume("resume.pdf", pdf_bytes("Python developer 3 years")))

    def test_image_or_blank_pdf_explains_ocr_and_paste(self):
        with self.assertRaisesRegex(ValueError, "OCR.*본문 붙여넣기"):
            extract_resume("scan.pdf", pdf_bytes())

    def test_encrypted_pdf_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "암호"):
            extract_resume("private.pdf", pdf_bytes("Private resume", password="secret"))

    def test_more_than_fifty_pages_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "50페이지"):
            extract_resume("long.pdf", pdf_bytes("Resume", pages=51))

    def test_invalid_pdf_and_wrong_signature(self):
        for data in (b"ordinary text", b"%PDF-1.7\nthis is broken"):
            with self.assertRaises(ValueError):
                extract_resume("resume.pdf", data)

    def test_pdf_text_limit(self):
        with self.assertRaisesRegex(ValueError, "50,000"):
            extract_resume("long.pdf", pdf_bytes("a" * (MAX_TEXT_CHARS + 1)))


class DocxImportTests(unittest.TestCase):
    def test_paragraphs_and_tables_keep_document_order(self):
        text = extract_resume("resume.docx", docx_bytes())
        self.assertEqual(text.splitlines(), [
            "김개발의 이력서", "기술\tPython", "경력\t데이터 분석 3년", "프로젝트 성과: 처리 시간 30% 단축"
        ])

    def test_wrong_signature_or_unrelated_zip_is_rejected(self):
        archive = BytesIO()
        with ZipFile(archive, "w") as zipped:
            zipped.writestr("unrelated.txt", "not a word file")
        for data in (b"not a zip", b"PK\x03\x04broken", archive.getvalue()):
            with self.assertRaises(ValueError):
                extract_resume("resume.docx", data)

    def test_decompression_limit_is_checked_before_document_load(self):
        archive = BytesIO()
        with ZipFile(archive, "w", compression=ZIP_DEFLATED) as zipped:
            zipped.writestr("word/document.xml", b"a" * (MAX_DOCX_UNCOMPRESSED_BYTES + 1))
        self.assertLess(len(archive.getvalue()), MAX_UPLOAD_BYTES)
        with self.assertRaisesRegex(ValueError, "30MB"):
            extract_resume("bomb.docx", archive.getvalue())

    def test_docx_text_limit(self):
        document = Document()
        document.add_paragraph("a" * (MAX_TEXT_CHARS + 1))
        archive = BytesIO()
        document.save(archive)
        with self.assertRaisesRegex(ValueError, "50,000"):
            extract_resume("long.docx", archive.getvalue())

    def test_malformed_document_xml_is_a_user_facing_error(self):
        source = ZipFile(BytesIO(docx_bytes()))
        archive = BytesIO()
        with source, ZipFile(archive, "w", compression=ZIP_DEFLATED) as zipped:
            for entry in source.infolist():
                zipped.writestr(entry, b"<broken" if entry.filename == "word/document.xml" else source.read(entry.filename))
        with self.assertRaisesRegex(ValueError, "본문 붙여넣기"):
            extract_resume("broken.docx", archive.getvalue())


if __name__ == "__main__":
    unittest.main()
